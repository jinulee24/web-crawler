"""도매꾹 베스트 상품 정기 수집 (월·수 08:00 스케줄에서 실행).

수집 대상 : https://www.domeggook.com/main/item/itemPopular.php?cat=<코드>
            대분류 카테고리 전부(페이지 네비에서 읽음) — 단 '가전/휴대폰/산업'(01_06)과 '전체'(01_00)는 제외
            카테고리당 한 페이지에 1~150위 전부 → 카테고리 하나에 요청 1회
출력      : ./<YYYYMMDD_HHMM>/
              도매꾹_베스트_전체_<YYYYMMDD_HHMM>.xlsx   시트 '전체' + 카테고리별 시트
              도매꾹_급상승10_NEW_<YYYYMMDD_HHMM>.xlsx  순위변동 ▲10 이상 + NEW (구분 열)
              raw_data.json, summary.json
            마지막 줄에 RUN_DIR=<경로> 를 찍는다 (스케줄 작업이 이 경로로 노션·메일 처리)

사다리 1단 plain_get (fingerprints/www_domeggook_com/profile.json 레시피 그대로)
  - 응답이 EUC-KR → resp.body.decode("cp949")
  - 아이템 루트는 div.gradeWrap 의 부모 li (topBig/topMiddle/topSmall 세 종류)

'순위변동' 열은 사이트 표기(▲ n / ▼ n / NEW / -) 그대로이고, 급상승 판정(▲ 10 이상)도 이 값으로 한다.
사이트는 비교 기간을 밝히지 않는다. 참고용으로 '전주순위'·'자체 전주대비' 열을 둔다 —
이 폴더에 쌓인 6~8일 전 실행 결과에서 같은 상품번호·카테고리의 순위를 찾아 직접 계산한 값.
"""
import glob
import json
import os
import re
import sys
from datetime import datetime, timedelta

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(BASE_DIR, "..", "..", ".."))
sys.path.insert(0, os.path.join(REPO, "scripts"))

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from scrapling.parser import Selector
from utils import RateLimiter, plain_get, detect_softblock, detect_pii, validate_values, setup_logger
from export_excel import export_to_excel

URL = "https://www.domeggook.com/main/item/itemPopular.php?cat={}"
EXCLUDE_CODES = {"01_00", "01_06"}          # 전체, 가전/휴대폰/산업
EXCLUDE_NAME_HINT = "가전"                   # 코드가 바뀌어도 이름으로 한 번 더 거른다
RISE_THRESHOLD = 10                          # ▲ 10 이상을 급상승으로 본다
FALLBACK_CATS = [
    ("01_01", "패션잡화/화장품"),
    ("01_02", "의류/언더웨어"),
    ("01_03", "출산/유아동/완구"),
    ("01_04", "가구/생활/취미"),
    ("01_05", "스포츠/건강/식품"),
]

logger = setup_logger("domeggook_weekly")
limiter = RateLimiter(delay=1.5, max_requests=12)


def fetch(cat):
    limiter.wait()
    resp = plain_get(URL.format(cat), timeout=30)
    html = resp.body.decode("cp949", "replace")
    return resp.status, html, Selector(html)


def read_categories(sel):
    cats = []
    for a in sel.css("a[href*='itemPopular.php?cat=01_']"):
        m = re.search(r"cat=(01_\d\d)", a.attrib.get("href", ""))
        name = " ".join(t.strip() for t in a.css("::text").getall() if t.strip())
        if not m or not name or m.group(1) in [c for c, _ in cats]:
            continue
        cats.append((m.group(1), name))
    cats = [(c, n) for c, n in cats if c not in EXCLUDE_CODES and EXCLUDE_NAME_HINT not in n]
    return cats if cats else FALLBACK_CATS


def parse(sel, cat_name):
    rows = []
    for item in sel.xpath("//div[@class='gradeWrap']/.."):
        rank_txt = (item.css("span.grade::text").get("") or "").strip()
        if not rank_txt.isdigit():
            continue
        href = item.css("div.title a::attr(href)").get("") or ""
        m = re.search(r"/(\d+)", href)
        pid = m.group(1) if m else ""
        price_txt = "".join(item.css("div.price::text").getall()).strip()
        price = int(re.sub(r"[^\d]", "", price_txt)) if re.search(r"\d", price_txt) else None
        note = ""
        if price is None:
            note = (item.css("div.setBuyTrueImgWrap img::attr(alt)").get("") or "").strip() or "가격 비공개"
        rows.append({
            "카테고리명": cat_name,
            "랭킹": int(rank_txt),
            "상품명": " ".join((item.css("div.title a::text").get("") or "").split()),
            "판매가격": price,
            "순위변동": " ".join((item.css("span.gradeUpDown::text").get("") or "").split()),
            "전주순위": None,
            "자체 전주대비": None,
            "상품번호": pid,
            "상품URL": f"https://www.domeggook.com/{pid}" if pid else "",
            "비고": note,
        })
    return rows


def jump(r):
    """사이트 순위변동의 상승 폭. ▲ 가 아니면 0."""
    if not r["순위변동"].startswith("▲"):
        return 0
    return int(re.sub(r"[^\d]", "", r["순위변동"]) or 0)


def attach_last_week(results, now):
    """6~8일 전 실행 중 7일 전에 가장 가까운 것을 찾아 전주순위·자체 전주대비를 붙인다."""
    best, best_gap = None, None
    for path in glob.glob(os.path.join(BASE_DIR, "*", "raw_data.json")):
        stamp = os.path.basename(os.path.dirname(path))
        try:
            t = datetime.strptime(stamp, "%Y%m%d_%H%M")
        except ValueError:
            continue
        gap = now - t
        if timedelta(days=6) <= gap <= timedelta(days=8):
            d = abs(gap - timedelta(days=7))
            if best_gap is None or d < best_gap:
                best, best_gap = path, d
    if not best:
        return None
    with open(best, encoding="utf-8") as f:
        prev = json.load(f)
    idx = {(r["카테고리명"], r["상품번호"]): r["랭킹"] for r in prev if r.get("상품번호")}
    for r in results:
        last = idx.get((r["카테고리명"], r["상품번호"]))
        r["전주순위"] = last
        r["자체 전주대비"] = (last - r["랭킹"]) if last else "신규"
    return os.path.basename(os.path.dirname(best))


def add_sheet(wb, rows, title):
    ws = wb.create_sheet(title[:31])
    headers = list(rows[0].keys()) if rows else ["결과"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h).font = Font(bold=True)
    for r_i, row in enumerate(rows, 2):
        for c, h in enumerate(headers, 1):
            ws.cell(row=r_i, column=c, value=row.get(h))
    for c, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r.get(h))) for r in rows if r.get(h) is not None])
        ws.column_dimensions[get_column_letter(c)].width = min(width + 2, 50)
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"


def main():
    now = datetime.now()
    stamp = now.strftime("%Y%m%d_%H%M")
    run_dir = os.path.join(BASE_DIR, stamp)

    try:
        status, html, sel = fetch("01_01")
    except Exception as e:
        logger.error(f"첫 페이지 실패: {e}")
        sys.exit(2)
    verdict = detect_softblock(html, status=status,
                               selector_hit=bool(sel.xpath("//div[@class='gradeWrap']/..")))
    if verdict["blocked"]:
        logger.error(f"소프트블록 감지 — {verdict['verdict']}: {verdict['signals']} — "
                     "사다리 B 는 통지 없이 가지 않는다. 중단")
        sys.exit(3)

    cats = read_categories(sel)
    logger.info("수집 카테고리: " + ", ".join(f"{c}={n}" for c, n in cats))

    results, failed, consecutive_errors = [], [], 0
    for cat, name in cats:
        try:
            if cat != "01_01":
                status, html, sel = fetch(cat)
            if status != 200:
                raise Exception(f"status {status}")
            rows = parse(sel, name)
            if not rows:
                raise Exception("핵심 셀렉터 미매칭")
            results.extend(rows)
            consecutive_errors = 0
            logger.info(f"[{name}] {len(rows)}건")
        except Exception as e:
            consecutive_errors += 1
            failed.append(name)
            logger.warning(f"[{name}] 실패({consecutive_errors}회): {e}")
            if consecutive_errors >= 5:
                break
            continue

    if not results:
        logger.error("0건 — 산출물을 만들지 않고 종료")
        sys.exit(4)

    os.makedirs(run_dir, exist_ok=True)
    prev_stamp = attach_last_week(results, now)

    risers = sorted([r for r in results if jump(r) >= RISE_THRESHOLD], key=lambda r: (-jump(r), r["랭킹"]))
    news = sorted([r for r in results if r["순위변동"] == "NEW"], key=lambda r: (r["카테고리명"], r["랭킹"]))
    highlight = [{"구분": f"급상승(▲{RISE_THRESHOLD}+)", **r} for r in risers] + [{"구분": "NEW", **r} for r in news]

    issues = validate_values(results, {
        "카테고리명": {"type": "str", "required": True, "allow_uniform": True},
        "랭킹":       {"type": "int", "required": True, "min": 1, "max": 1000},
        "상품명":     {"type": "str", "required": True, "max_empty_ratio": 0.0},
    })
    issues += validate_values([r for r in results if r["판매가격"] is not None], {
        "판매가격": {"type": "int", "required": True, "min": 1, "max": 100_000_000},
    })
    pii = detect_pii(results)

    with open(os.path.join(run_dir, "raw_data.json"), "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)

    # 파일 1 — 전체 목록 (시트 '전체' + 카테고리별 시트)
    all_xlsx = export_to_excel(results, os.path.join(run_dir, f"도매꾹_베스트_전체_{stamp}.xlsx"), sheet_name="전체")
    wb = load_workbook(all_xlsx)
    for _, name in cats:
        rows = [r for r in results if r["카테고리명"] == name]
        if rows:
            add_sheet(wb, rows, name.replace("/", "·"))
    wb.save(all_xlsx)

    # 파일 2 — 급상승(▲10+) + NEW
    hl_path = os.path.join(run_dir, f"도매꾹_급상승{RISE_THRESHOLD}_NEW_{stamp}.xlsx")
    if highlight:
        hl_xlsx = export_to_excel(highlight, hl_path, sheet_name="급상승·NEW")
    else:
        hl_xlsx = export_to_excel([{"결과": "해당 상품 없음"}], hl_path, sheet_name="급상승·NEW")

    per_cat = {}
    for r in results:
        c = per_cat.setdefault(r["카테고리명"], {"total": 0, "rise10": 0, "new": 0, "price_hidden": 0})
        c["total"] += 1
        c["rise10"] += jump(r) >= RISE_THRESHOLD
        c["new"] += r["순위변동"] == "NEW"
        c["price_hidden"] += r["판매가격"] is None

    keys = ("카테고리명", "랭킹", "상품명", "판매가격", "순위변동", "상품URL")
    cand_keys = keys + ("전주순위",)
    # Brief 후보 — 상승폭과 현재 순위를 함께 본다 (상위권에서 크게 오른 상품이 앞으로)
    def notable(r):
        return jump(r) * (1 + (150 - r["랭킹"]) / 150)
    gift_words = ("선물세트", "추석", "명절", "설날")
    summary_extra = {
        "brief_candidates_rising": [{k: r[k] for k in cand_keys} for r in
                                    sorted(risers, key=lambda r: -notable(r))[:30]],
        "brief_candidates_new": [{k: r[k] for k in cand_keys} for r in
                                 sorted(news, key=lambda r: r["랭킹"])[:30]],
        "gift_keyword_count": sum(1 for r in highlight if any(w in r["상품명"] for w in gift_words)),
    }
    summary = {
        "run_at": now.isoformat(timespec="minutes"),
        "run_dir": run_dir,
        "all_xlsx": all_xlsx,
        "highlight_xlsx": hl_xlsx,
        "rise_threshold": RISE_THRESHOLD,
        "categories": [n for _, n in cats],
        "failed_categories": failed,
        "total": len(results),
        "rise10_total": len(risers),
        "new_total": len(news),
        "highlight_total": len(highlight),
        "per_category": per_cat,
        "top_movers": [{k: r[k] for k in keys} for r in risers[:10]],
        "new_top": [{k: r[k] for k in keys} for r in sorted(news, key=lambda r: r["랭킹"])[:10]],
        "last_week_run": prev_stamp,
        "validation_issues": issues,
        "pii_warnings": pii,
        "requests": limiter.request_count,
        **summary_extra,
    }
    with open(os.path.join(run_dir, "summary.json"), "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)

    logger.info(f"총 {len(results)}건 · 급상승(▲{RISE_THRESHOLD}+) {len(risers)}건 · NEW {len(news)}건 · "
                f"실패 카테고리 {failed or '없음'} · 요청 {limiter.request_count}회")
    if issues:
        logger.warning("값 검증 경고:\n" + "\n".join(issues))
    print(f"RUN_DIR={run_dir}")


if __name__ == "__main__":
    main()
